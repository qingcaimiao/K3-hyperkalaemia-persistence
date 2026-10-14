# -*- coding: utf-8 -*-
"""
md2docx.py —— Markdown → Word 转换器
支持: 标题(#/##/###) / 粗体 / 斜体 / 表格 / 有序与无序列表 / 引用块 / 水平线 / 插图
插图语法:  ![](绝对路径.png "宽度英寸")   —— 必须独占一行
用法:  python md2docx.py 输入.md 输出.docx [西文字体] [中文字体]
"""
import sys, os, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run_font(run, ascii_font='Arial', ea_font='宋体', size=None, bold=None, italic=None):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), ascii_font)
    rf.set(qn('w:hAnsi'), ascii_font)
    rf.set(qn('w:eastAsia'), ea_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`|\^[^^\n]+\^)')


def _add_sup(par, tok, ascii_font, ea_font, base_size, bold_all):
    """把 ^1^ / ^2,3^ 渲染为真正的上标 run（投稿稿不得出现裸^ 标记）。"""
    sup = tok[1:-1]
    if sup.startswith('*') and sup.endswith('*'):   # ^*对应^ 的星号
        sup = sup[1:-1]
    r = par.add_run(sup)
    set_run_font(r, ascii_font, ea_font, base_size, bold_all)
    r.font.superscript = True
    return r


def _unescape(t):
    """去除 markdown 转义反斜杠：\\* -> *   \\^ -> ^"""
    return t.replace(r'\*', '*').replace(r'\^', '^')


def add_inline(par, text, ascii_font, ea_font, base_size, bold_all=False):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('^') and tok.endswith('^') and len(tok) > 2:
            _add_sup(par, tok, ascii_font, ea_font, base_size, bold_all)
        elif tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            r = par.add_run(_unescape(tok[2:-2])); set_run_font(r, ascii_font, ea_font, base_size, True)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = par.add_run(_unescape(tok[1:-1])); set_run_font(r, ascii_font, ea_font, base_size, bold_all, True)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = par.add_run(tok[1:-1]); set_run_font(r, 'Consolas', ea_font, base_size - 0.5, bold_all)
        else:
            # 处理 markdown 转义符：\* -> *  \^ -> ^  （投稿稿不得出现反斜杠）
            t = tok.replace(r'\*', '*').replace(r'\^', '^')
            r = par.add_run(t); set_run_font(r, ascii_font, ea_font, base_size, bold_all)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)


def convert(md_path, out_path, ascii_font='Arial', ea_font='宋体'):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = ascii_font
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), ea_font)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(0.9)

    i = 0
    n_img = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()

        # 空行
        if not s:
            i += 1; continue

        # 插图
        m = re.match(r'^!\[\]\((.+?)(?:\s+"([\d.]+)")?\)\s*$', s)
        if m:
            p = m.group(1); w = float(m.group(2)) if m.group(2) else 6.2
            # 相对路径一律以 md 文件所在目录为基准解析（支持可移植的图片引用）
            if not os.path.isabs(p):
                p = os.path.join(os.path.dirname(os.path.abspath(md_path)), p)
            if os.path.exists(p):
                par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(p, width=Inches(w))
                par.paragraph_format.keep_with_next = True   # 图与其图注不得被分页拆开
                n_img += 1
            else:
                par = doc.add_paragraph(); add_inline(par, f'[缺图: {p}]', ascii_font, ea_font, 9)
            i += 1; continue

        # 水平线
        if re.match(r'^-{3,}$', s):
            doc.add_paragraph()
            i += 1; continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            lvl = len(m.group(1)); txt = m.group(2).strip()
            par = doc.add_paragraph()
            sizes = {1: 16, 2: 13.5, 3: 12, 4: 11, 5: 10.5, 6: 10.5}
            add_inline(par, txt, ascii_font, ea_font, sizes.get(lvl, 11), bold_all=True)
            if lvl <= 2:
                par.paragraph_format.space_before = Pt(10)
                par.paragraph_format.space_after = Pt(6)
            else:
                par.paragraph_format.space_before = Pt(8)
                par.paragraph_format.space_after = Pt(4)
            i += 1; continue

        # 表格
        if s.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i + 1].strip()):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip()); i += 1
            header = [c.strip() for c in block[0].strip('|').split('|')]
            body = [[c.strip() for c in r.strip('|').split('|')] for r in block[2:]]
            ncol = len(header)
            if doc.paragraphs:                               # 表题与表格同页，不跨页分离
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
            t = doc.add_table(rows=1, cols=ncol)
            t.style = 'Table Grid'
            for j, h in enumerate(header):
                c = t.rows[0].cells[j]; c.text = ''
                add_inline(c.paragraphs[0], h, ascii_font, ea_font, 9, bold_all=True)
                shade(c, 'EDEDED')
            for r in body:
                cells = t.add_row().cells
                for j in range(ncol):
                    v = r[j] if j < len(r) else ''
                    cells[j].text = ''
                    add_inline(cells[j].paragraphs[0], v, ascii_font, ea_font, 9)
            # 表格跨页时重复标题行（MDPI 审稿最常挑的排版问题）
            trPr = t.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
            doc.add_paragraph()
            continue

        # 引用块
        if s.startswith('>'):
            txt = s.lstrip('>').strip()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.3)
            add_inline(par, txt, ascii_font, ea_font, 10)
            for r in par.runs:
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1; continue

        # 列表
        m = re.match(r'^(\s*)([-*+]|\d+\.)\s+(.*)$', ln)
        if m:
            indent = len(m.group(1)) // 2
            txt = m.group(3)
            marker = m.group(2)
            if re.match(r'^\d+\.', marker):
                # 有序列表：保留 markdown 的显式编号为纯文本 + 悬挂缩进。
                # 不可改用 Word 的 'List Number' 自动编号样式——多个独立列表共享同一
                # 编号定义，Word 会连续累加（署名 1,2 之后参考文献就从 3 开始），
                # 且编号不在文本层，检索与投稿解析都会出错。
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Inches(0.5 + 0.25 * indent)
                par.paragraph_format.first_line_indent = Inches(-0.3)
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.line_spacing = 1.15
                add_inline(par, f'{marker} {txt}', ascii_font, ea_font, 10.5)
            else:
                try:
                    par = doc.add_paragraph(style='List Bullet')
                except Exception:
                    par = doc.add_paragraph()
                par.paragraph_format.left_indent = Inches(0.25 + 0.25 * indent)
                par.paragraph_format.space_after = Pt(2)
                add_inline(par, txt, ascii_font, ea_font, 10.5)
            i += 1; continue

        # 普通段落
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(5)
        par.paragraph_format.line_spacing = 1.25
        add_inline(par, s, ascii_font, ea_font, 10.5)
        i += 1

    # 文档属性：默认作者会显示为转换工具名（python-docx），投稿件须改为真实作者
    _title = 'Manuscript'
    for _line in open(md_path, encoding='utf-8'):
        if _line.startswith('# '):
            _title = _line[2:].strip()
            break
    _cp = doc.core_properties
    _cp.author = 'Qingcai Miao'
    _cp.last_modified_by = 'Qingcai Miao'
    _cp.title = _title
    _cp.comments = ''
    _cp.category = ''
    _cp.subject = ''

    doc.save(out_path)
    return n_img


if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    a = sys.argv[3] if len(sys.argv) > 3 else 'Arial'
    e = sys.argv[4] if len(sys.argv) > 4 else '宋体'
    n = convert(src, dst, a, e)
    print(f'OK  {dst}   (嵌入图片 {n} 张)')
