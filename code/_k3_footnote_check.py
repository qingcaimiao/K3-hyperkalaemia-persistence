# -*- coding: utf-8 -*-
"""逐句复核表注：核对每条表注里的数字是否都能在审计 CSV 中找到"""
import pandas as pd, re, zipfile
from docx import Document

BS = chr(92)
a = pd.concat([pd.read_csv('out/k3_audit_numbers.csv'), pd.read_csv('out/k3_audit_numbers2.csv')], ignore_index=True)
vals = set()


def norm(s):
    try:
        f = float(str(s).replace(',', ''))
        return ('%g' % f)
    except Exception:
        return str(s)


for col in ['value', 'note', 'item']:
    for v in a[col].dropna().astype(str):
        for tok in re.findall(r'\d[\d,]*(?:\.\d+)?', v):
            vals.add(norm(tok))
# 稿件本身允许出现的常量：阈值、窗口、参考文献号
CONST = {'5.5', '24', '72', '7', '6', '12', '2', '4', '1', '3', '5.9', '6.5', '1.96', '9'}

files = ['out/manuscript_K3_EN.docx', 'out/稿件_K3_中文版.docx']
for f in files:
    z = zipfile.ZipFile(f); xml = z.read('word/document.xml').decode('utf-8')
    doc = Document(f)
    paras = [p.text for p in doc.paragraphs]
    print('=' * 90)
    print(f, '| 反斜杠:', xml.count(BS), '| 段落:', len(paras), '| 表:', len(doc.tables))
    # 提取表注：以 "Table N." / "表 N." 之后、或表格下方的说明段
    for i, t in enumerate(paras):
        if re.match(r'^(Table \d|表 \d)\b', t.strip()) or t.strip().startswith(('Baseline potassium', 'Confirmed episodes include',
                                                                                'The coefficient of variation', 'Confirmation was defined',
                                                                                'In this table', 'Coverage is', '基线血钾', '已确认事件包含',
                                                                                '变异系数', '确认定义为', '本表四种状态', '"已复查"')):
            print('\n[表注] %s' % t)
            nums = re.findall(r'\d[\d,]*(?:\.\d+)?', t)
            miss = [n for n in nums if norm(n) not in vals and norm(n) not in CONST]
            print('   数字:', nums)
            print('   未在审计结果中出现的数字:', miss if miss else '无')
