# -*- coding: utf-8 -*-
"""K3 终检：中英数字一致性 + Word 反斜杠 + 非常规码点 + 图引用 + 表编号"""
import io, re, sys, unicodedata
from docx import Document
import zipfile

en = io.open('out/manuscript_K3.md', encoding='utf-8').read()
cn = io.open('out/稿件_K3_中文版.md', encoding='utf-8').read()

NUMS = ['5.82', '3.46', '1,182', '997', '224', '234', '86', '100', '757', '493', '189', '169', '150', '235',
        '3.72', '1.71', '0.93', '0.59', '0.42', '0.35', '0.74', '0.82', '73.4', '64.7', '78.3', '69.3',
        '19.0', '23.5', '7.3', '10.0', '1,032', '762', '4.85', '5.10', '4.60', '4.70',
        '20,322', '28,805', '23,429', '41,917', '98', '682', '2,128', '3,626', '881', '8,804',
        '19,140', '27,808', '18,227', '23,087', '7,427', '10,571', '17,585', '22,602',
        '412', '502', '102', '138', '39', '40', '89', '75', '245', '325', '108', '106',
        '11.8', '17.4', '22.8', '23.1', '51.2', '53.0', '4.4', '5.8', '7.8', '13.1', '18.6', '15.7',
        '41.0', '34.9', '11.2', '16.7', '20.0', '15.1', '14.9', '27.5', '14.6', '14.7', '14.2',
        '1.56', '4.41', '11.87', '2.77', '4.52', '8.05', '1.91', '2.98', '3.47', '2.75', '3.02', '5.61',
        '0.99', '2.45', '2.32', '7.57', '5.59', '20.49', '1.82', '3.93', '1.85', '6.30', '4.16', '15.15',
        '1.29', '2.74', '1.90', '4.51', '1.41', '6.55', '1.60', '2.31', '8.79',
        '9.71', '9.40', '0.410', '0.386', '1.14', '1.07', '78.8', '60.1', '37.8', '82.8', '67.7', '46.5',
        '14.7', '19.2', '38', '83', '22', '31', '64.0', '64.8', '46.3', '51.1',
        '32.3', '33.9', '10.7', '1.8', '1.9', '39.0', '54.5', '29.3', '31.4', '1,698', '1,236',
        '52.0', '42.0', '15.4', '20.0', '35.9', '32.0', '76.7', '60.2', '87.3', '76.4',
        '137', '151', '4.5', '2.8', '615', '419', '567', '578', '61.2', '64.5', '94.2', '86.3',
        '375', '31', '36.3', '4.1', '27.5', '18.4', '29.0', '30.8', '1.60', '1.17', '2.10', '1.30', '0.72', '2.38',
        '32.9', '35.5', '11.7', '12.2', '24.2', '39.7', '22.7', '36.0', '82.0', '81.5', '72.2', '80.1',
        '32.0', '32.8', '40', '46.5', '69.0', '9.62', '19.28', '8.10', '13.39', '3.44', '19.50', '5.79', '39.32',
        '1.77', '16.43', '1.33', '28.97', '8.04', '3.31', '16.91', '2.04', '6.19', '1.35', '2.17', '3.32', '7.69',
        '2.12', '3.01', '3.10', '4.58', '7.28', '2.75', '15.07',
        '2,629', '1,998', '643', '326', '32.2', '16.3', '2,316', '1,422', '592', '309', '41.6', '21.7',
        '567', '498', '83', '22', '432', '357', '113', '34', '31.7', '9.5', '308', '280', '12.1', '3.2',
        '205', '164', '51', '16', '31.1', '9.8', '6.7', '1.2', '1.6', '1.7', '1.4']

bad = []
for n in NUMS:
    a, b = en.count(n), cn.count(n)
    if a != b:
        bad.append((n, a, b))
print('=== 中英数字计数不一致 (%d) ===' % len(bad))
for n, a, b in bad:
    print('   %-8s EN %d / CN %d' % (n, a, b))

# 图引用
print('\n=== 图正文引用 ===')
for i in range(1, 7):
    print('  Figure %d: EN %d 次 | 图 %d: CN %d 次' % (i, len(re.findall(r'Figure %d(?![0-9])' % i, en)),
                                                 i, len(re.findall(r'图 %d(?![0-9])' % i, cn))))
# 表编号
print('\n=== 表编号顺序 ===')
print('  EN:', re.findall(r'### Table (\d+)\.', en))
print('  CN:', re.findall(r'### 表 (\d+)\.', cn))

# 写 Word 并检查
PY = None
for cand in ['C:/Users/miaoqingcai/.workbuddy/binaries/python/envs/default/Scripts/python.exe']:
    PY = cand
import subprocess
for src, dst, font in [('out/manuscript_K3.md', 'out/manuscript_K3_EN.docx', 'Arial'),
                       ('out/稿件_K3_中文版.md', 'out/稿件_K3_中文版.docx', None)]:
    cmd = [PY, 'md2docx.py', src, dst] + ([font] if font else [])
    r = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace')
    print('\n[md2docx]', dst, '→', (r.stdout or '').strip().splitlines()[-1] if r.stdout else r.stderr[-200:])

for f in ['out/manuscript_K3_EN.docx', 'out/稿件_K3_中文版.docx']:
    z = zipfile.ZipFile(f); xml = z.read('word/document.xml').decode('utf-8')
    d = Document(f)
    txt = '\n'.join(p.text for p in d.paragraphs) + '\n'.join(
        c.text for t in d.tables for r in t.rows for c in r.cells)
    allow = set(range(0x20, 0x7f)) | {0x09, 0x0a}
    extra = set('—–‘’“”±×°²³≥≤−χ²Δκ→←·√τσλ̄')
    badc = {}
    for ch in set(txt):
        o = ord(ch)
        if o in allow or ch in extra or '\u4e00' <= ch <= '\u9fff' or ch in '，。、；：（）【】〔〕％':
            continue
        badc[ch] = (hex(o), unicodedata.name(ch, '?'))
    print('\n--- %s ---' % f)
    print('  反斜杠: %d | 段落: %d | 表格: %d | 图片: %d' % (
        xml.count('\\'), len(d.paragraphs), len(d.tables),
        sum(1 for r in d.part.rels.values() if 'image' in r.reltype)))
    print('  非常规字符:', badc if badc else '无')
    print('  占位符:', '发现' if re.search(r'XXX|TODO|待填|\[insert', txt) else '无')
